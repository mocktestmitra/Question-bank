"""Parse transcripts in data/<Subject>/*.txt, validate, and build Excel workbooks.

usage: python3 scripts/build.py            # all subjects
       python3 scripts/build.py Civil      # one subject

Transcript format (see header of data/Civil/1-92.txt):
  @chapter N Name
  ## QNO | star | tag=ES 99 | p=4 | hw=b | fig=PAGE:x0,y0,x1,y1 | check=reason | fix=reason
  question lines...
  #a option text (continuation lines allowed)
  #b ... #c ... #d ...
  #note text printed after the options
Optional answer keys (Appendix A) in data/<Subject>/answers.txt:   "CH QNO LETTER" per line
Optional explanations (Appendix B) in data/<Subject>/explanations.txt:
  "## CH QNO" followed by explanation lines.
"""
import os, re, sys, glob, subprocess, tempfile
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OPTS = "abcd"


def parse(path):
    chapters, qs = {}, []
    ch = None
    cur = None
    field = None
    pdf = os.path.relpath(os.path.join(ROOT, os.path.basename(os.path.dirname(path)),
                                       os.path.basename(path)[:-4] + ".pdf"), ROOT)
    for ln, raw in enumerate(open(path, encoding="utf-8"), 1):
        line = raw.rstrip("\n")
        if line.startswith("# ") or line == "#":
            continue
        if line.startswith("@chapter "):
            _, n, name = line.split(" ", 2)
            ch = int(n)
            chapters[ch] = name.strip()
            continue
        if line.startswith("## "):
            parts = [p.strip() for p in line[3:].split("|")]
            meta = {"star": False}
            for p in parts[1:]:
                if p == "star":
                    meta["star"] = True
                elif "=" in p:
                    k, v = p.split("=", 1)
                    meta[k.strip()] = v.strip()
                else:
                    raise SystemExit(f"{path}:{ln}: bad meta '{p}'")
            cur = {"ch": ch, "qno": parts[0], "q": [], "opts": {}, "note": [], "src": pdf, "line": ln, **meta}
            qs.append(cur)
            field = cur["q"]
            continue
        m = re.match(r"^#([a-e]|note) ?(.*)$", line)
        if m and cur is not None:
            k, txt = m.groups()
            if k == "note":
                field = cur["note"]
            else:
                if k in cur["opts"]:
                    raise SystemExit(f"{path}:{ln}: duplicate option {k} in Q{cur['qno']}")
                cur["opts"][k] = []
                field = cur["opts"][k]
            field.append(txt)
            continue
        if line.strip() == "":
            continue
        if cur is None:
            raise SystemExit(f"{path}:{ln}: text outside a question")
        field.append(line)
    for q in qs:
        q["q"] = "\n".join(q["q"]).strip()
        q["note"] = "\n".join(q["note"]).strip()
        q["opts"] = {k: "\n".join(v).strip() for k, v in q["opts"].items()}
    return chapters, qs


def load_keys(subj_dir):
    keys, expl = {}, {}
    a = os.path.join(subj_dir, "answers.txt")
    if os.path.exists(a):
        for line in open(a, encoding="utf-8"):
            if line.strip() and not line.startswith("#"):
                c, n, l = line.split()
                keys[(int(c), n)] = l.lower()
    b = os.path.join(subj_dir, "explanations.txt")
    if os.path.exists(b):
        cur = None
        for line in open(b, encoding="utf-8"):
            if line.startswith("## "):
                c, n = line[3:].split()[:2]
                cur = (int(c), n)
                expl[cur] = []
            elif cur and not line.startswith("# "):
                expl[cur].append(line.rstrip("\n"))
        expl = {k: "\n".join(v).strip() for k, v in expl.items()}
    return keys, expl


def validate(chapters, qs, keys, expl):
    issues = []
    by_ch = {}
    for q in qs:
        by_ch.setdefault(q["ch"], []).append(q)
    for ch, lst in sorted(by_ch.items()):
        nums = [int(q["qno"]) for q in lst]
        seen = set()
        for n in nums:
            if n in seen:
                issues.append((ch, n, f"duplicate Q number {n}"))
            seen.add(n)
        for n in range(1, max(nums) + 1):
            if n not in seen:
                issues.append((ch, n, f"missing Q number {n}"))
        if nums != sorted(nums):
            issues.append((ch, None, "Q numbers out of order"))
    for q in qs:
        flags = []
        if not q["q"]:
            flags.append("empty question")
        for o in OPTS:
            if not q["opts"].get(o):
                flags.append(f"option {o} missing/empty")
        hw = q.get("hw", "")
        if hw and hw not in OPTS:
            flags.append(f"handwritten mark '{hw}' not a-d")
        if q.get("check"):
            flags.append(q["check"])
        k = keys.get((q["ch"], q["qno"]))
        q["key"] = k or ""
        q["expl"] = expl.get((q["ch"], q["qno"]), "")
        if not k:
            q["match"] = "NO_KEY"
        elif not hw:
            q["match"] = "NO_HANDWRITTEN"
        elif k == hw:
            q["match"] = "OK"
        else:
            q["match"] = "MISMATCH"
            flags.append(f"Appendix A key '{k}' != handwritten '{hw}'")
        if keys and q["star"] and not q["expl"]:
            flags.append("starred but no Appendix B explanation")
        q["flags"] = flags
        q["conf"] = "CHECK" if flags else "HIGH"
    return issues


def crop_figure(q, outdir):
    page, box = q["fig"].split(":")
    x0, y0, x1, y1 = map(float, box.split(","))
    os.makedirs(outdir, exist_ok=True)
    out = os.path.join(outdir, f"ch{q['ch']:02d}_q{int(q['qno']):03d}.png")
    if not os.path.exists(out):
        with tempfile.TemporaryDirectory() as d:
            subprocess.run(["pdftoppm", "-f", page, "-l", page, "-r", "200", "-png",
                            os.path.join(ROOT, q["src"]), f"{d}/p"], check=True)
            im = Image.open(os.path.join(d, os.listdir(d)[0]))
            w, h = im.size
            im.crop((int(x0 * w), int(y0 * h), int(x1 * w), int(y1 * h))).save(out)
    return out


HEAD = ["Subject", "Chapter No", "Chapter", "Q No", "Starred", "Question",
        "Option a", "Option b", "Option c", "Option d",
        "Correct Option", "Correct Option Text", "Explanation", "Exam Tag",
        "Handwritten Mark", "Key Match", "Has Figure", "Figure",
        "Source File", "PDF Page", "Confidence", "Notes"]
WIDTH = {"Question": 60, "Option a": 28, "Option b": 28, "Option c": 28, "Option d": 28,
         "Correct Option Text": 28, "Explanation": 50, "Chapter": 22, "Notes": 40, "Figure": 26,
         "Source File": 18}


def row(subj, chapters, q):
    question = q["q"] + ("\n" + q["note"] if q["note"] else "")
    notes = list(q["flags"])
    if q.get("fix"):
        notes.append(q["fix"])
    if q["note"]:
        notes.append("Last line(s) of question printed after the options in the book.")
    return [subj, q["ch"], chapters[q["ch"]], int(q["qno"]), "Y" if q["star"] else "",
            question, *(q["opts"].get(o, "") for o in OPTS),
            q["key"], q["opts"].get(q["key"], "") if q["key"] else "", q["expl"], q.get("tag", ""),
            q.get("hw", ""), q["match"], "Y" if q.get("fig") else "", "",
            q["src"], q.get("p", ""), q["conf"], "; ".join(notes)]


def write_sheet(ws, rows, figs):
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r, f in zip(rows, figs):
        ws.append(r)
        i = ws.max_row
        for c in ws[i]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if r[HEAD.index("Confidence")] == "CHECK":
            ws.cell(i, HEAD.index("Confidence") + 1).fill = PatternFill("solid", fgColor="FFC7CE")
        if f:
            img = XLImage(f)
            scale = min(1.0, 180 / img.width, 140 / img.height)
            img.width, img.height = int(img.width * scale), int(img.height * scale)
            ws.add_image(img, f"{get_column_letter(HEAD.index('Figure') + 1)}{i}")
            ws.row_dimensions[i].height = max(img.height * 0.75 + 4, 15)
    for j, h in enumerate(HEAD, 1):
        ws.column_dimensions[get_column_letter(j)].width = WIDTH.get(h, 11)
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = ws.dimensions


def build(subj):
    sdir = os.path.join(ROOT, "data", subj)
    chapters, qs = {}, []
    for f in sorted(glob.glob(os.path.join(sdir, "*.txt"))):
        if os.path.basename(f) in ("answers.txt", "explanations.txt"):
            continue
        c, q = parse(f)
        chapters.update(c)
        qs += q
    keys, expl = load_keys(sdir)
    issues = validate(chapters, qs, keys, expl)
    outdir = os.path.join(ROOT, "output", subj)
    figdir = os.path.join(outdir, "figures")
    os.makedirs(outdir, exist_ok=True)
    qs.sort(key=lambda q: (q["ch"], int(q["qno"])))

    wb = Workbook()
    idx = wb.active
    idx.title = "Index"
    idx.append(["Chapter No", "Chapter", "Questions", "Starred", "With Key", "Handwritten Marks", "Flagged (CHECK)", "Sheet"])
    for c in idx[1]:
        c.font = Font(bold=True)
    allrows, allfigs = [], []
    for ch in sorted(chapters):
        lst = [q for q in qs if q["ch"] == ch]
        rows = [row(subj, chapters, q) for q in lst]
        figs = [crop_figure(q, figdir) if q.get("fig") else None for q in lst]
        name = f"{ch:02d} {chapters[ch]}"[:31]
        write_sheet(wb.create_sheet(name), rows, figs)
        allrows += rows
        allfigs += figs
        idx.append([ch, chapters[ch], len(lst), sum(q["star"] for q in lst), sum(bool(q["key"]) for q in lst),
                    sum(bool(q.get("hw")) for q in lst), sum(q["conf"] == "CHECK" for q in lst), name])
    write_sheet(wb.create_sheet("All"), allrows, allfigs)
    for j, w in enumerate([11, 40, 11, 9, 9, 18, 15, 30], 1):
        idx.column_dimensions[get_column_letter(j)].width = w
    out = os.path.join(outdir, f"{subj}_QuestionBank.xlsx")
    wb.save(out)

    print(f"{subj}: {len(qs)} questions, {len(chapters)} chapters -> {os.path.relpath(out, ROOT)}")
    for ch, n, msg in issues:
        print(f"  ISSUE ch{ch} Q{n}: {msg}")
    for q in qs:
        if q["conf"] == "CHECK":
            print(f"  CHECK ch{q['ch']} Q{q['qno']} (p{q.get('p')}): {'; '.join(q['flags'])}")


if __name__ == "__main__":
    subs = sys.argv[1:] or sorted(d for d in os.listdir(os.path.join(ROOT, "data"))
                                  if os.path.isdir(os.path.join(ROOT, "data", d)))
    for s in subs:
        build(s)
